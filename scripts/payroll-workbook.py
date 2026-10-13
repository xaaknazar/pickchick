#!/usr/bin/env python3
"""Read the owner's two-sheet timesheet. Private JSON only; never modifies the xlsx.

python payroll-workbook.py workbook.xlsx private.json
Decimal hours are preserved as integer milliseconds. Blank cells remain unknown.
Every payroll amount retains its source cell, cached value and formula for review.
"""
import collections
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
import re
import sys
import openpyxl


def extract(path):
    path = Path(path)
    raw = path.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    formulas = openpyxl.load_workbook(path, data_only=False)
    values = openpyxl.load_workbook(path, data_only=True)
    if set(values.sheetnames) != {'График', 'Табель'}:
        raise ValueError('Expected График and Табель sheets')
    graph, payroll = values['График'], values['Табель']
    title = str(graph['A1'].value)
    months = ['январь','февраль','март','апрель','май','июнь','июль','август','сентябрь','октябрь','ноябрь','декабрь']
    match = re.search(r'(' + '|'.join(months) + r')\s+(20\d{2})', title, re.I)
    if not match:
        raise ValueError('Unrecognized period')
    month = f'{match[2]}-{months.index(match[1].lower())+1:02}-01'
    if str(graph['AH2'].value).strip() != 'Общее кол.во Факт часов':
        raise ValueError('Actual-hours column required; do not infer plans from facts')
    norm = lambda s: ' '.join(str(s).split()).casefold()
    decimal = lambda v: Decimal(str(v))
    def amount(v):
        if v is None: return None
        if not isinstance(v, (int, float)): raise ValueError('Non-numeric money')
        return str(int((decimal(v)*100).quantize(Decimal('1'), rounding=ROUND_HALF_UP)))
    rows = {}
    for r in range(4, min(graph.max_row,10000)+1):
        if graph.cell(r,2).value is None: continue
        name = ' '.join(str(graph.cell(r,2).value).split())
        if norm(name) in rows: raise ValueError('Duplicate employee name; explicit mapping required')
        days = []
        for c in range(4,34):
            value = graph.cell(r,c).value
            if value is None: continue
            if not isinstance(value,(int,float)) or isinstance(value,bool) or not 0 < value <= 24:
                raise ValueError('Unrecognized day code or duration at '+graph.cell(r,c).coordinate)
            ms = decimal(value)*3600000
            if ms != int(ms): raise ValueError('Submillisecond duration')
            days.append({'date':month[:8]+f'{c-3:02}', 'paid_duration_ms':int(ms), 'source':{'sha256':digest,'cell':'График!'+graph.cell(r,c).coordinate,'hours':str(value)}})
        cached = graph.cell(r,34).value
        if cached is None or abs(decimal(cached)*3600000-sum(d['paid_duration_ms'] for d in days)) > 1:
            raise ValueError('Daily hours disagree with total at row '+str(r))
        rows[norm(name)] = {'name':name,'position':str(graph.cell(r,3).value or '').strip(),'graph_row':r,'days':days,'payroll':None,'rate_minor':None}
    for r in range(4,min(payroll.max_row,10000)+1):
        name=payroll.cell(r,2).value
        if name is None: continue
        person=rows.get(norm(name))
        if not person:
            person={'name':' '.join(str(name).split()),'position':str(payroll.cell(r,3).value or '').strip(),'graph_row':None,'days':[],'payroll':None,'rate_minor':None}
            rows[norm(name)]=person
        if person['payroll'] is not None: raise ValueError('Ambiguous payroll employee')
        gross,net=amount(payroll.cell(r,7).value),amount(payroll.cell(r,21).value)
        if gross is None or net is None: raise ValueError('Missing cached payroll result')
        items=[]
        for c in range(8,20):
            value=payroll.cell(r,c).value
            if value is None or value==0: continue
            minor=amount(value)
            if int(minor)<0: raise ValueError('Negative adjustment requires explicit review')
            items.append({'name':str(payroll.cell(2,c).value).strip(),'effect':'addition' if c==8 else 'deduction','amount_minor':minor})
        expected=int(gross)+sum(int(i['amount_minor'])*(1 if i['effect']=='addition' else -1) for i in items)
        issues=[]
        if expected!=int(net):issues.append('Итог Excel отличается от суммы базы, доплат и удержаний. Требуется сверка.')
        formula=formulas['Табель'].cell(r,7).value
        if isinstance(formula,str) and formula.startswith('=') and formula!=f'=D{r}*E{r}':
            issues.append('В исходной базе есть ручная формула или несколько ставок. Даты смены ставок не указаны.')
        cells=[]
        for c in range(3,26):
            if c in (22,23): continue # duplicate name and blank signature are not financial facts
            cell=payroll.cell(r,c); formula=formulas['Табель'].cell(r,c)
            cells.append({'cell':cell.coordinate,'label':str(payroll.cell(2,c).value or ''),'value':'' if cell.value is None else str(cell.value),'formula':formula.value if formula.data_type=='f' else ''})
        person['rate_minor']=amount(payroll.cell(r,4).value)
        person['payroll']={'base_override_minor':gross,'items':items,'source':{'sha256':digest,'sheet':'Табель','row':r,'gross_minor':gross,'net_minor':net,'cells':cells,'issues':issues}}
    return {'version':1,'sha256':digest,'month':month,'employees':list(rows.values()),'summary':{'employees':len(rows),'days':sum(len(p['days']) for p in rows.values()),'duration_ms':sum(d['paid_duration_ms'] for p in rows.values() for d in p['days']),'payroll_rows':sum(p['payroll'] is not None for p in rows.values()),'source_net_minor':str(sum(int(p['payroll']['source']['net_minor']) for p in rows.values() if p['payroll']))}}

if __name__=='__main__':
    import os
    os.umask(0o077)
    result=extract(sys.argv[1])
    out=Path(sys.argv[2]);out.write_text(json.dumps(result,ensure_ascii=False,indent=2))
    out.chmod(0o600)
    print(json.dumps({'sha256':result['sha256'],'month':result['month'],'summary':result['summary']},ensure_ascii=False))
